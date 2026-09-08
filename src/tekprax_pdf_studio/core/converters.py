from __future__ import annotations

from pathlib import Path
import os
import shutil
import subprocess
import tempfile

import fitz
from pdf2docx import Converter
from docx import Document
from docx.shared import Inches


def find_libreoffice() -> str | None:
    for name in ("libreoffice", "soffice"):
        found = shutil.which(name)
        if found:
            return found

    candidates: list[Path] = []
    if os.name == "nt":
        for env_name in ("PROGRAMFILES", "PROGRAMFILES(X86)"):
            base = os.environ.get(env_name)
            if base:
                candidates.append(Path(base) / "LibreOffice" / "program" / "soffice.exe")
    elif os.uname().sysname == "Darwin":
        candidates.append(Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"))

    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    return None


def office_to_pdf(input_file: str | Path, output_file: str | Path) -> Path:
    soffice = find_libreoffice()
    if not soffice:
        raise RuntimeError(
            "LibreOffice não encontrado. Instale o LibreOffice para converter DOC/DOCX/PPT/PPTX em PDF."
        )

    input_file = Path(input_file).resolve()
    output_file = Path(output_file).resolve()
    output_file.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="tekprax_lo_") as tmp:
        tmpdir = Path(tmp)
        proc = subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(tmpdir),
                str(input_file),
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=180,
        )
        if proc.returncode != 0:
            raise RuntimeError(proc.stderr.strip() or proc.stdout.strip() or "Falha na conversão com LibreOffice.")

        generated = tmpdir / f"{input_file.stem}.pdf"
        if not generated.exists():
            pdfs = list(tmpdir.glob("*.pdf"))
            if not pdfs:
                raise RuntimeError("O LibreOffice não gerou o PDF esperado.")
            generated = pdfs[0]

        shutil.move(str(generated), str(output_file))

    return output_file


def _configure_tesseract() -> bool:
    try:
        import pytesseract
    except Exception:
        return False

    if os.name == "nt":
        common = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
        if common.exists():
            pytesseract.pytesseract.tesseract_cmd = str(common)

    try:
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _has_searchable_text(pdf_path: str | Path, min_chars: int = 30) -> bool:
    doc = fitz.open(str(pdf_path))
    try:
        chars = 0
        for page in doc:
            chars += len(page.get_text("text").strip())
            if chars >= min_chars:
                return True
        return False
    finally:
        doc.close()


def _ocr_pdf_to_docx(input_pdf: str | Path, output_docx: str | Path) -> Path:
    import pytesseract
    from PIL import Image

    if not _configure_tesseract():
        raise RuntimeError("Tesseract OCR não está instalado ou não foi encontrado.")

    doc = fitz.open(str(input_pdf))
    out = Document()
    try:
        for index, page in enumerate(doc):
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            try:
                text = pytesseract.image_to_string(img, lang="por+eng")
            except Exception:
                text = pytesseract.image_to_string(img, lang="eng")

            out.add_heading(f"Página {index + 1}", level=2)
            out.add_paragraph(text.strip() or "[Nenhum texto reconhecido]")
            if index < len(doc) - 1:
                out.add_page_break()
    finally:
        doc.close()

    output_docx = Path(output_docx)
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    out.save(str(output_docx))
    return output_docx


def pdf_to_docx(input_pdf: str | Path, output_docx: str | Path, use_ocr_fallback: bool = True) -> tuple[Path, bool]:
    output_docx = Path(output_docx)
    output_docx.parent.mkdir(parents=True, exist_ok=True)

    if use_ocr_fallback and not _has_searchable_text(input_pdf):
        _ocr_pdf_to_docx(input_pdf, output_docx)
        return output_docx, True

    cv = Converter(str(input_pdf))
    try:
        cv.convert(str(output_docx), start=0, end=None)
    finally:
        cv.close()
    return output_docx, False
